from __future__ import annotations

import csv
from pathlib import Path
from typing import Dict, Sequence


PARTICLE_COLUMNS = [
    "particle_id",
    "area_um2",
    "perimeter_um",
    "perimeter_crofton_um",
    "perimeter_subpixel_um",
    "q_value",
    "roundness",
    "roundness_method",
    "major_axis_um",
    "minor_axis_um",
    "axis_ratio",
    "hole_area_um2",
    "hole_ratio",
    "equivalent_diameter_um",
    "feret_diameter_um",
    "bbox_max_diameter_um",
    "statistical_diameter_um",
    "statistical_diameter_method",
    "touches_image_border",
    "is_spherical",
    "spherical_rule",
    "spherical_roundness",
    "spherical_roundness_method",
    "spherical_roundness_fallback",
    "is_hollow",
    "is_agglomerate",
    "agglomerate_group_id",
    "agglomerate_group_size",
    "agglomerate_contact_degree",
    "agglomerate_max_contact_ratio",
    "agglomerate_max_overlap_ratio",
    "agglomerate_min_gap_um",
    "score",
    "centroid_x",
    "centroid_y",
]

RENAME_MAP = {
    "particle_id": "颗粒ID",
    "area_um2": "面积A(um^2)",
    "perimeter_um": "周长P(um)",
    "perimeter_crofton_um": "Crofton周长(um)",
    "perimeter_subpixel_um": "亚像素周长(um)",
    "q_value": "球形度Q",
    "roundness": "圆形度sqrt(Q)",
    "roundness_method": "圆形度周长方法",
    "major_axis_um": "长轴(um)",
    "minor_axis_um": "短轴(um)",
    "axis_ratio": "轴比Lmajor/Lminor",
    "hole_area_um2": "孔洞面积(um^2)",
    "hole_ratio": "孔洞占比",
    "equivalent_diameter_um": "等效直径(um)",
    "feret_diameter_um": "最大Feret径(um)",
    "bbox_max_diameter_um": "外接框最大径(um)",
    "statistical_diameter_um": "统计粒径(um)",
    "statistical_diameter_method": "粒径统计方法",
    "touches_image_border": "是否接触图像边界",
    "is_spherical": "是否球形颗粒",
    "spherical_rule": "球形判定规则",
    "spherical_roundness": "球形判定圆形度",
    "spherical_roundness_method": "球形判定圆形度方法",
    "spherical_roundness_fallback": "球形判定特征是否回退",
    "is_hollow": "是否空心粉",
    "is_agglomerate": "是否团聚体",
    "agglomerate_group_id": "团聚体Group ID",
    "agglomerate_group_size": "团聚组颗粒数",
    "agglomerate_contact_degree": "强接触邻居数",
    "agglomerate_max_contact_ratio": "最大接触弧占比",
    "agglomerate_max_overlap_ratio": "最大掩膜重叠占比",
    "agglomerate_min_gap_um": "最小接触间隙(um)",
    "score": "分割置信度",
    "centroid_x": "中心X(px)",
    "centroid_y": "中心Y(px)",
}


def _particle_rows(features: Sequence[Dict]):
    keep_columns = [col for col in PARTICLE_COLUMNS if any(col in item for item in features)]
    headers = [RENAME_MAP.get(col, col) for col in keep_columns]
    rows = [[item.get(col, "") for col in keep_columns] for item in features]
    return headers, rows


def export_excel_report(
    output_path: str | Path,
    features: Sequence[Dict],
    stats: Dict,
) -> None:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    headers, rows = _particle_rows(features)
    stats_headers = [
        "总颗粒数N",
        "平均球形度Q",
        "平均圆形度sqrt(Q)",
        "空心粉数n_hollow",
        "空心粉率K",
        "团聚体数N_agglom_group",
        "团聚颗粒数N_agglom_particle",
        "团聚率P_agglom",
        "团聚面积率P_area",
        "球形颗粒数",
        "球形颗粒率S",
        "团聚体group",
        "团聚体pair",
        "团聚候选pair数",
        "通过强接触pair数",
        "最终入组pair数",
        "团聚组诊断",
        "请求物理容差(um)",
        "实际物理容差(um)",
        "实际像素容差(px)",
        "标尺(um/px)",
        "容差来源",
        "最小接触弧占比",
        "最小掩膜重叠占比",
        "最小成组颗粒数",
        "团聚判定口径",
    ]
    stats_headers.extend(
        [
            "粒径口径",
            "有效粒径统计颗粒数",
            "边界颗粒数",
            "平均粒径(um)",
            "Dmin(um)",
            "D5(um)",
            "D10(um)",
            "D50(um)",
            "D90(um)",
            "D95(um)",
            "Dmax(um)",
            "圆形度周长口径",
            "球形判定规则",
            "球形圆形度阈值C",
            "球形轴比阈值",
        ]
    )
    stats_headers.extend(
        [
            "球形判定圆形度方法",
            "球形判定特征回退颗粒数",
        ]
    )
    stats_row = [
        stats["total_particles"],
        stats["mean_sphericity_q_text"],
        stats.get("mean_roundness_text", ""),
        stats["hollow_particles"],
        stats["hollow_rate_text"],
        stats.get("agglomerate_group_count", 0),
        stats["agglomerate_particles"],
        stats["agglomerate_rate_text"],
        stats["agglomerate_area_rate_text"],
        stats["spherical_particles"],
        stats["sphericity_rate_s_text"],
        str(stats.get("agglomerate_groups", [])),
        str(stats.get("agglomerate_pairs", [])),
        stats.get("agglomerate_candidate_pair_count", 0),
        stats.get("agglomerate_accepted_pair_count", 0),
        stats.get("agglomerate_group_pair_count", 0),
        str(stats.get("agglomerate_group_details", [])),
        stats.get("agglomerate_tolerance_um", ""),
        stats.get("agglomerate_effective_tolerance_um", ""),
        stats.get("agglomerate_tolerance_px", ""),
        stats.get("agglomerate_pixel_size_um", ""),
        stats.get("agglomerate_tolerance_source", ""),
        stats.get("agglomerate_min_contact_ratio", ""),
        stats.get("agglomerate_min_overlap_ratio", ""),
        stats.get("agglomerate_min_group_size", ""),
        stats.get("agglomerate_definition", ""),
    ]
    stats_row.extend(
        [
            stats.get("diameter_method", ""),
            stats.get("size_statistics_particle_count", 0),
            stats.get("border_particle_count", 0),
            stats.get("diameter_mean_um", 0.0),
            stats.get("diameter_min_um", 0.0),
            stats.get("diameter_d5_um", 0.0),
            stats.get("diameter_d10_um", 0.0),
            stats.get("diameter_d50_um", 0.0),
            stats.get("diameter_d90_um", 0.0),
            stats.get("diameter_d95_um", 0.0),
            stats.get("diameter_max_um", 0.0),
            stats.get("roundness_method", ""),
            stats.get("spherical_rule", ""),
            stats.get("spherical_roundness_threshold", ""),
            stats.get("spherical_axis_ratio_threshold", ""),
        ]
    )
    stats_row.extend(
        [
            stats.get("spherical_roundness_method", ""),
            stats.get("spherical_roundness_fallback_particles", 0),
        ]
    )

    try:
        from openpyxl import Workbook

        workbook = Workbook()
        stats_sheet = workbook.active
        stats_sheet.title = "统计结果"
        stats_sheet.append(stats_headers)
        stats_sheet.append(stats_row)

        particle_sheet = workbook.create_sheet("颗粒级特征")
        particle_sheet.append(headers)
        for row in rows:
            particle_sheet.append(row)

        evidence_sheet = workbook.create_sheet("团聚判定证据")
        evidence_headers = [
            "颗粒A",
            "颗粒B",
            "间隙(px)",
            "间隙(um)",
            "接触弧占比",
            "掩膜重叠占比",
            "中心距/半径和",
            "是否通过",
            "判定原因",
            "人工真值(待填写)",
        ]
        evidence_sheet.append(evidence_headers)
        for item in stats.get("agglomerate_pair_evidence", []):
            evidence_sheet.append(
                [
                    item.get("particle_a", ""),
                    item.get("particle_b", ""),
                    item.get("gap_px", ""),
                    item.get("gap_um", ""),
                    item.get("contact_ratio", ""),
                    item.get("overlap_ratio", ""),
                    item.get("center_distance_ratio", ""),
                    item.get("accepted", ""),
                    item.get("reason", ""),
                    "",
                ]
            )

        workbook.save(output_path)
    except Exception:
        # Excel writer unavailable时，至少保证CSV报告可用。
        output_path.with_suffix(".stats.csv").write_text(
            ",".join(stats_headers) + "\n" + ",".join(str(v) for v in stats_row) + "\n",
            encoding="utf-8-sig",
        )

    csv_path = output_path.with_suffix(".particles.csv")
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)

    evidence_csv_path = output_path.with_suffix(".agglomerate_pairs.csv")
    with open(evidence_csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(
            [
                "particle_a",
                "particle_b",
                "gap_px",
                "gap_um",
                "contact_ratio",
                "overlap_ratio",
                "center_distance_ratio",
                "accepted",
                "reason",
                "manual_truth",
            ]
        )
        for item in stats.get("agglomerate_pair_evidence", []):
            writer.writerow(
                [
                    item.get("particle_a", ""),
                    item.get("particle_b", ""),
                    item.get("gap_px", ""),
                    item.get("gap_um", ""),
                    item.get("contact_ratio", ""),
                    item.get("overlap_ratio", ""),
                    item.get("center_distance_ratio", ""),
                    item.get("accepted", ""),
                    item.get("reason", ""),
                    "",
                ]
            )
