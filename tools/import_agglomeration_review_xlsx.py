from __future__ import annotations

import argparse
import csv
import json
import math
from pathlib import Path
from typing import Dict, Iterable

from openpyxl import load_workbook


VALID_LABELS = {"", "0", "1", "-1"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Copy manual_truth labels from an Excel review workbook back onto "
            "the original CSV without trusting Excel-converted identifier cells."
        )
    )
    parser.add_argument(
        "--template-csv",
        default="",
        help=(
            "Optional original CSV. When it is unavailable, the workbook rows "
            "are restored directly and batch_id is recovered from source_csv."
        ),
    )
    parser.add_argument("--filled-xlsx", required=True)
    parser.add_argument("--output-csv", required=True)
    parser.add_argument("--sheet", default="")
    return parser.parse_args()


def normalize_source_row(value: object) -> str:
    text = str(value or "").strip()
    try:
        parsed = float(text)
    except ValueError:
        return text
    if math.isfinite(parsed) and parsed.is_integer():
        return str(int(parsed))
    return text


def review_key(row: Dict[str, object]) -> str:
    source_csv = str(row.get("source_csv") or "").strip().replace("\\", "/")
    source_row = normalize_source_row(row.get("source_row"))
    if not source_csv or not source_row:
        raise ValueError(
            "Every review row must contain source_csv and source_row."
        )
    return f"{Path(source_csv).name.casefold()}#{source_row}"


def normalize_label(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    if text.endswith(".0"):
        try:
            number = float(text)
        except ValueError:
            pass
        else:
            if number.is_integer():
                text = str(int(number))
    if text not in VALID_LABELS:
        raise ValueError(
            f"Unsupported manual_truth value {value!r}; use 1, 0, -1, or blank."
        )
    return text


def batch_id_from_source_csv(value: object) -> str:
    name = Path(str(value or "").strip().replace("\\", "/")).name
    for suffix in (
        "_merged.agglomerate_pairs.csv",
        ".agglomerate_pairs.csv",
        "_merged.csv",
        ".csv",
    ):
        if name.casefold().endswith(suffix.casefold()):
            return name[: -len(suffix)]
    return name


def workbook_rows(path: Path, sheet_name: str = "") -> Iterable[Dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = (
            workbook[sheet_name]
            if sheet_name
            else workbook[workbook.sheetnames[0]]
        )
        values = worksheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values)]
        required = {"source_csv", "source_row", "manual_truth"}
        missing = sorted(required - set(headers))
        if missing:
            raise ValueError(
                f"Review workbook is missing columns: {', '.join(missing)}"
            )
        for values_row in values:
            yield dict(zip(headers, values_row))
    finally:
        workbook.close()


def main() -> None:
    args = parse_args()
    template_path = Path(args.template_csv) if args.template_csv else None
    workbook_path = Path(args.filled_xlsx)
    output_path = Path(args.output_csv)

    workbook_data = list(workbook_rows(workbook_path, str(args.sheet or "")))
    labels: Dict[str, str] = {}
    for row in workbook_data:
        key = review_key(row)
        label = normalize_label(row.get("manual_truth"))
        previous = labels.get(key)
        if previous is not None and previous != label:
            raise ValueError(
                f"Conflicting labels for {key}: {previous!r} vs {label!r}"
            )
        labels[key] = label

    use_template = bool(template_path and template_path.is_file())
    if use_template:
        with template_path.open("r", newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            if not reader.fieldnames:
                raise ValueError(f"Template CSV has no header: {template_path}")
            rows = list(reader)
            fieldnames = list(reader.fieldnames)
    else:
        rows = [dict(row) for row in workbook_data]
        fieldnames = list(rows[0]) if rows else []
        for row in rows:
            row["source_row"] = normalize_source_row(row.get("source_row"))
            row["manual_truth"] = normalize_label(row.get("manual_truth"))
            if "batch_id" in row:
                row["batch_id"] = batch_id_from_source_csv(row.get("source_csv"))

    template_keys = {review_key(row) for row in rows}
    unknown_keys = sorted(set(labels) - template_keys)
    if unknown_keys:
        raise ValueError(
            "Workbook contains rows absent from the template CSV: "
            + ", ".join(unknown_keys[:5])
        )

    for row in rows:
        key = review_key(row)
        if key in labels:
            row["manual_truth"] = labels[key]

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    counts = {
        label: sum(
            normalize_label(row.get("manual_truth")) == label for row in rows
        )
        for label in ("1", "0", "-1", "")
    }
    summary = {
        "template_csv": (
            str(template_path.resolve()) if use_template and template_path else None
        ),
        "restored_from_workbook": not use_template,
        "filled_xlsx": str(workbook_path.resolve()),
        "output_csv": str(output_path.resolve()),
        "rows": len(rows),
        "positive": counts["1"],
        "negative": counts["0"],
        "invalid_mask": counts["-1"],
        "blank": counts[""],
    }
    summary_path = output_path.with_suffix(".summary.json")
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
