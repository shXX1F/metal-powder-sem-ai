from __future__ import annotations

import csv
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

import cv2
import numpy as np

from metal_powder_sem_ai.classify_stat import (
    DEFAULT_SPHERICAL_ROUNDNESS_METHOD,
    DEFAULT_SPHERICAL_ROUNDNESS_THRESHOLD,
    classify_particles,
    reference_percentile,
    summarize_size_distribution,
)
from metal_powder_sem_ai.feature_extract import extract_particle_features
from metal_powder_sem_ai.report import export_excel_report
from metal_powder_sem_ai.segment import InstanceMask, bbox_from_mask, crop_mask_to_bbox, mask_to_contour
from metal_powder_sem_ai.visualize import class_name_and_color


def circle_instance(center=(30, 30), radius=10, shape=(80, 80), particle_id=1):
    mask = np.zeros(shape, dtype=np.uint8)
    cv2.circle(mask, center, radius, 1, thickness=-1)
    contour = mask_to_contour(mask)
    assert contour is not None
    bbox = bbox_from_mask(mask)
    return InstanceMask(
        particle_id=particle_id,
        mask=crop_mask_to_bbox(mask, bbox),
        contour=contour,
        bbox_xyxy=bbox,
        score=0.9,
    )


class CalibratedStatisticsTests(unittest.TestCase):
    def test_circle_feret_and_roundness(self):
        feature = extract_particle_features(
            circle_instance(),
            pixel_size_um=0.5,
            image_shape=(80, 80),
        )
        self.assertAlmostEqual(feature["feret_diameter_px"], 20.0, delta=1.0)
        self.assertGreater(feature["roundness_crofton"], 0.90)
        self.assertLessEqual(feature["roundness_crofton"], 1.0)
        for key in (
            "roundness_crofton2",
            "roundness_crofton_close3",
            "roundness_crofton_open3",
            "roundness_crofton_smooth3",
        ):
            self.assertGreater(feature[key], 0.0)
            self.assertLessEqual(feature[key], 1.0)
        self.assertFalse(feature["touches_image_border"])

    def test_border_flag(self):
        feature = extract_particle_features(
            circle_instance(center=(8, 30), radius=8),
            pixel_size_um=1.0,
            image_shape=(80, 80),
        )
        self.assertTrue(feature["touches_image_border"])
        self.assertFalse(feature["valid_for_size_statistics"])

    def test_reference_percentile_uses_lower_rank(self):
        self.assertEqual(reference_percentile([1.0, 2.0, 3.0, 4.0], 50), 2.0)
        self.assertEqual(reference_percentile([1.0, 2.0, 3.0, 4.0], 90), 3.0)

    def test_size_summary_uses_selected_diameter(self):
        features = [
            {"feret_diameter_um": 1.0, "equivalent_diameter_um": 0.5},
            {"feret_diameter_um": 2.0, "equivalent_diameter_um": 1.0},
            {"feret_diameter_um": 3.0, "equivalent_diameter_um": 1.5},
            {"feret_diameter_um": 4.0, "equivalent_diameter_um": 2.0},
        ]
        summary = summarize_size_distribution(features, method="feret_max")
        self.assertEqual(summary["diameter_d50_um"], 2.0)
        self.assertEqual(summary["diameter_mean_um"], 2.5)

    def test_classification_preserves_calibrated_method(self):
        feature = extract_particle_features(
            circle_instance(),
            pixel_size_um=0.5,
            image_shape=(80, 80),
        )
        classified, stats = classify_particles(
            [feature],
            diameter_method="feret_max",
            roundness_method="crofton",
        )
        self.assertEqual(stats["diameter_method"], "feret_max")
        self.assertEqual(stats["roundness_method"], "crofton")
        self.assertEqual(classified[0]["roundness"], feature["roundness_crofton"])

    def test_spherical_rule_uses_calibrated_roundness(self):
        feature = extract_particle_features(
            circle_instance(),
            pixel_size_um=0.5,
            image_shape=(80, 80),
        )
        feature["axis_ratio"] = 5.0
        classified, stats = classify_particles(
            [feature],
            roundness_method="crofton",
            spherical_roundness_threshold=0.90,
            spherical_rule="roundness",
        )
        self.assertTrue(classified[0]["is_spherical"])
        self.assertEqual(stats["spherical_rule"], "roundness")

    def test_spherical_classifier_uses_calibrated_blend_without_changing_report(self):
        feature = extract_particle_features(
            circle_instance(),
            pixel_size_um=0.5,
            image_shape=(80, 80),
        )
        classified, stats = classify_particles(
            [feature],
            spherical_roundness_method=DEFAULT_SPHERICAL_ROUNDNESS_METHOD,
            spherical_roundness_threshold=DEFAULT_SPHERICAL_ROUNDNESS_THRESHOLD,
        )
        expected = (
            0.75 * feature["roundness_crofton_open3"]
            + 0.25 * feature["roundness_crofton_smooth3"]
        )
        self.assertAlmostEqual(classified[0]["roundness"], feature["roundness_crofton"])
        self.assertAlmostEqual(classified[0]["spherical_roundness"], expected)
        self.assertEqual(
            classified[0]["spherical_roundness_method"],
            DEFAULT_SPHERICAL_ROUNDNESS_METHOD,
        )
        self.assertFalse(classified[0]["spherical_roundness_fallback"])
        self.assertEqual(stats["spherical_roundness_fallback_particles"], 0)

    def test_spherical_classifier_falls_back_for_legacy_features(self):
        feature = extract_particle_features(
            circle_instance(),
            pixel_size_um=0.5,
            image_shape=(80, 80),
        )
        feature.pop("roundness_crofton_open3")
        classified, stats = classify_particles(
            [feature],
            spherical_roundness_method=DEFAULT_SPHERICAL_ROUNDNESS_METHOD,
            spherical_roundness_threshold=DEFAULT_SPHERICAL_ROUNDNESS_THRESHOLD,
        )
        self.assertEqual(
            classified[0]["spherical_roundness"],
            feature["roundness_crofton"],
        )
        self.assertTrue(classified[0]["spherical_roundness_fallback"])
        self.assertEqual(stats["spherical_roundness_fallback_particles"], 1)

    def test_agglomerate_tolerance_is_scale_aware(self):
        feature_half_um = extract_particle_features(
            circle_instance(),
            pixel_size_um=0.5,
            image_shape=(80, 80),
        )
        _, stats_half_um = classify_particles(
            [feature_half_um],
            agglomerate_tolerance_um=0.75,
        )
        feature_one_um = extract_particle_features(
            circle_instance(),
            pixel_size_um=1.0,
            image_shape=(80, 80),
        )
        _, stats_one_um = classify_particles(
            [feature_one_um],
            agglomerate_tolerance_um=0.75,
        )
        self.assertEqual(stats_half_um["agglomerate_tolerance_px"], 2)
        self.assertEqual(stats_one_um["agglomerate_tolerance_px"], 1)
        self.assertEqual(stats_half_um["agglomerate_tolerance_source"], "physical_um")
        self.assertEqual(stats_one_um["agglomerate_tolerance_source"], "physical_um")

    def test_strong_contact_component_is_agglomerate(self):
        instances = [
            circle_instance(center=(30, 40), radius=12, shape=(100, 100), particle_id=1),
            circle_instance(center=(48, 40), radius=12, shape=(100, 100), particle_id=2),
            circle_instance(center=(39, 55), radius=12, shape=(100, 100), particle_id=3),
        ]
        features = [
            extract_particle_features(instance, pixel_size_um=0.5, image_shape=(100, 100))
            for instance in instances
        ]
        classified, stats = classify_particles(
            features,
            masks=[instance.mask for instance in instances],
            agglomerate_tolerance_um=0.5,
            min_agglomerate_group_size=3,
            agglomerate_min_contact_ratio=0.12,
            agglomerate_min_overlap_ratio=0.03,
        )
        self.assertEqual(stats["agglomerate_group_count"], 1)
        self.assertEqual(stats["agglomerate_particles"], 3)
        self.assertTrue(all(item["is_agglomerate"] for item in classified))
        self.assertGreaterEqual(stats["agglomerate_group_pair_count"], 2)

    def test_point_contacts_are_not_automatically_agglomerate(self):
        instances = [
            circle_instance(center=(20, 40), radius=10, shape=(100, 100), particle_id=1),
            circle_instance(center=(40, 40), radius=10, shape=(100, 100), particle_id=2),
            circle_instance(center=(60, 40), radius=10, shape=(100, 100), particle_id=3),
        ]
        features = [
            extract_particle_features(instance, pixel_size_um=0.5, image_shape=(100, 100))
            for instance in instances
        ]
        classified, stats = classify_particles(
            features,
            masks=[instance.mask for instance in instances],
            agglomerate_tolerance_um=0.5,
            min_agglomerate_group_size=3,
            agglomerate_min_contact_ratio=0.12,
            agglomerate_min_overlap_ratio=0.03,
        )
        self.assertEqual(stats["agglomerate_group_count"], 0)
        self.assertEqual(stats["agglomerate_particles"], 0)
        self.assertFalse(any(item["is_agglomerate"] for item in classified))

    def test_report_exports_calibration_metadata(self):
        feature = extract_particle_features(
            circle_instance(),
            pixel_size_um=0.5,
            image_shape=(80, 80),
        )
        classified, stats = classify_particles([feature])
        with TemporaryDirectory() as directory:
            output = Path(directory) / "report.xlsx"
            export_excel_report(output, classified, stats)
            fallback = output.with_suffix(".stats.csv")
            if fallback.exists():
                with fallback.open("r", newline="", encoding="utf-8-sig") as handle:
                    headers, values = list(csv.reader(handle))[:2]
            else:
                from openpyxl import load_workbook

                workbook = load_workbook(output, read_only=False, data_only=True)
                statistics = workbook["统计结果"]
                worksheet_rows = statistics.iter_rows()
                headers = [cell.value for cell in next(worksheet_rows)]
                values = [cell.value for cell in next(worksheet_rows)]
                workbook.close()
            row = dict(zip(headers, values))
            self.assertEqual(row["粒径口径"], "feret_max")
            self.assertEqual(row["圆形度周长口径"], "crofton")
            self.assertEqual(row["球形判定规则"], "roundness")
            self.assertAlmostEqual(
                float(row["球形圆形度阈值C"]),
                DEFAULT_SPHERICAL_ROUNDNESS_THRESHOLD,
            )
            evidence_csv = output.with_suffix(".agglomerate_pairs.csv")
            self.assertTrue(evidence_csv.exists())

    def test_visualization_class_codes_are_ascii(self):
        for feature in (
            {"is_agglomerate": True, "is_spherical": False},
            {"is_agglomerate": False, "is_spherical": True},
            {"is_agglomerate": False, "is_spherical": False},
        ):
            class_code, _ = class_name_and_color(feature)
            self.assertTrue(class_code.isascii())


if __name__ == "__main__":
    unittest.main()
